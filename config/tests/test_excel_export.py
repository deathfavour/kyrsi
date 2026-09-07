from io import BytesIO

from openpyxl import load_workbook

from config.excel_export import export_queryset_to_xlsx


def test_export_queryset_to_xlsx_writes_header_and_rows():
    rows = ["Alice", "Bob"]
    columns = [("Name", lambda r: r), ("Length", lambda r: len(r))]

    response = export_queryset_to_xlsx(rows, columns, "test.xlsx")

    assert response["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response["Content-Disposition"] == 'attachment; filename="test.xlsx"'

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == ["Name", "Length"]
    assert [cell.value for cell in sheet[2]] == ["Alice", 5]
    assert [cell.value for cell in sheet[3]] == ["Bob", 3]


def test_export_queryset_to_xlsx_empty_queryset_has_header_only():
    response = export_queryset_to_xlsx([], [("Name", lambda r: r)], "empty.xlsx")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    assert sheet.max_row == 1

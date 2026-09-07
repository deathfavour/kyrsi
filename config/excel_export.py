import datetime

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _excel_safe(value):
    """openpyxl rejects timezone-aware datetimes outright (Excel has no tz
    concept) — USE_TZ=True means every DateTimeField value is tz-aware, so
    every export would raise without this."""
    if isinstance(value, datetime.datetime) and value.tzinfo is not None:
        return timezone.localtime(value).replace(tzinfo=None)
    return value


def export_queryset_to_xlsx(queryset, columns, filename):
    """Generic admin-action helper: dumps `columns` (list of (header, accessor)
    pairs, accessor is a callable(obj) -> value) for every row in `queryset`
    into a single-sheet .xlsx and returns it as a download response.

    Column layout is a reasonable default per model (see CLAUDE.md "Excel-экспорт" —
    exact format is still open pending accounting sign-off), not a
    customer-confirmed template.
    """
    workbook = Workbook()
    sheet = workbook.active

    for col_index, (header, _accessor) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=col_index, value=header)
        cell.font = Font(bold=True)

    for row_index, obj in enumerate(queryset, start=2):
        for col_index, (_header, accessor) in enumerate(columns, start=1):
            sheet.cell(row=row_index, column=col_index, value=_excel_safe(accessor(obj)))

    for col_index in range(1, len(columns) + 1):
        sheet.column_dimensions[get_column_letter(col_index)].width = 22

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response
